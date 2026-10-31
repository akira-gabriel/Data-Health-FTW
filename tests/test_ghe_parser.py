import openpyxl
import pytest

from data_health_ftw.ghe.parser import OUTPUT_COLUMNS, parse_asdr_workbook

COUNTRIES = [("Brazil", "BRA", 1), ("Chile", "CHL", 2)]


def add_asdr_sheet(wb, year, rows):
    """Mimic the WHO layout: header rows 7-9, population row 10, cause rows from 11."""
    ws = wb.create_sheet(f"ASDR {year}")
    ws.cell(row=7, column=1, value="Sex")
    for i, (name, iso, flag) in enumerate(COUNTRIES):
        ws.cell(row=7, column=8 + i, value=name)
        ws.cell(row=8, column=8 + i, value=iso)
        ws.cell(row=9, column=8 + i, value=flag)
    ws.cell(row=10, column=1, value="Persons")
    ws.cell(row=10, column=4, value="Population ('000) (2)")
    for i in range(len(COUNTRIES)):
        ws.cell(row=10, column=8 + i, value=1000.0)
    for r, (sex, code, labels, values) in enumerate(rows, start=11):
        ws.cell(row=r, column=1, value=sex)
        ws.cell(row=r, column=2, value=code)
        for offset, label in enumerate(labels):  # label cells sit in columns C..F
            ws.cell(row=r, column=3 + offset, value=label)
        for i, value in enumerate(values):
            ws.cell(row=r, column=8 + i, value=value)


@pytest.fixture
def workbook_path(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Notes"
    add_asdr_sheet(
        wb,
        2023,
        [
            ("Persons", 1100, (None, "H.", "Cardiovascular diseases"), (300.5, 250.0)),
            ("Persons", 1180, (None, None, "1.", "Chronic obstructive pulmonary disease"), (50.0, None)),
            ("Males", 1100, (None, "H.", "Cardiovascular diseases"), (400.0, 350.0)),
        ],
    )
    add_asdr_sheet(wb, 2000, [("Persons", 1100, (None, "H.", "Cardiovascular diseases"), (500.0, 450.0))])
    path = tmp_path / "asdr.xlsx"
    wb.save(path)
    return path


def test_parse_returns_expected_columns_and_years(workbook_path):
    df = parse_asdr_workbook(workbook_path)
    assert list(df.columns) == OUTPUT_COLUMNS
    assert sorted(df["year"].unique()) == [2000, 2023]


def test_parse_skips_notes_sheet_and_population_row(workbook_path):
    df = parse_asdr_workbook(workbook_path)
    assert not df["cause"].str.contains("Population").any()
    assert set(df["ghe_code"]) == {1100, 1180}


def test_parse_reads_country_metadata_and_values(workbook_path):
    df = parse_asdr_workbook(workbook_path)
    row = df[(df.country_code == "BRA") & (df.year == 2023) & (df.sex == "Persons") & (df.ghe_code == 1100)]
    assert len(row) == 1
    assert row.iloc[0]["asdr_per_100k"] == 300.5
    assert row.iloc[0]["country_name"] == "Brazil"
    assert row.iloc[0]["country_flag"] == 1


def test_parse_uses_last_filled_label_cell_as_cause(workbook_path):
    df = parse_asdr_workbook(workbook_path)
    causes = dict(zip(df["ghe_code"], df["cause"]))
    assert causes[1100] == "Cardiovascular diseases"
    assert causes[1180] == "Chronic obstructive pulmonary disease"


def test_parse_keeps_sex_blocks_apart(workbook_path):
    df = parse_asdr_workbook(workbook_path)
    males = df[(df.sex == "Males") & (df.country_code == "BRA") & (df.year == 2023)]
    assert males.iloc[0]["asdr_per_100k"] == 400.0


def test_parse_keeps_missing_values_as_null(workbook_path):
    df = parse_asdr_workbook(workbook_path)
    row = df[(df.country_code == "CHL") & (df.ghe_code == 1180)]
    assert row["asdr_per_100k"].isna().all()


def test_parse_has_unique_keys(workbook_path):
    df = parse_asdr_workbook(workbook_path)
    assert not df.duplicated(["country_code", "year", "sex", "ghe_code"]).any()


def test_parse_raises_when_no_asdr_sheets(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Notes"
    path = tmp_path / "empty.xlsx"
    wb.save(path)
    with pytest.raises(ValueError):
        parse_asdr_workbook(path)